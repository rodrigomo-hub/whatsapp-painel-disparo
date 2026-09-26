"""Leitura limitada de CSV/XLSX; nenhum cadastro ou envio nesta etapa."""
import csv
import io
import json
import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook

MAX_BYTES = 10 * 1024 * 1024
COLUNAS = {
    'telefone': 'telefone', 'whatsapp': 'telefone', 'phone': 'telefone', 'celular': 'telefone',
    'cpf': 'cpf', 'nome': 'nome', 'name': 'nome', 'nome_completo': 'nome',
    'nascimento': 'nascimento', 'data_de_nascimento': 'nascimento', 'data_nascimento': 'nascimento', 'birth_date': 'nascimento',
    'email': 'email', 'e_mail': 'email', 'produto': 'produto', 'product': 'produto',
    'campanha': 'campanha', 'campaign': 'campanha', 'variaveis_template': 'variaveis_template',
    **{f'variavel_{i}': f'variavel_{i}' for i in range(1, 11)},
}


def cabecalho(valor):
    texto = unicodedata.normalize('NFKD', str(valor or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '_', texto).strip('_')


def ler(nome, conteudo):
    if not conteudo or len(conteudo) > MAX_BYTES:
        raise HTTPException(400, 'Arquivo vazio ou maior que 10 MB')
    tipo = Path(nome).suffix.lower()
    workbook = None
    try:
        if tipo == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
                if sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                    raise HTTPException(400, 'Planilha descompactada excede 40 MB')
            workbook = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=False, keep_links=False)
            sheet = workbook.worksheets[0]
            if sheet.max_column and sheet.max_column > 100:
                raise HTTPException(400, 'Use ate 100 colunas na primeira aba')
            linhas = (tuple(c.value for c in row) for row in sheet.iter_rows())
        elif tipo == '.csv':
            try:
                texto = conteudo.decode('utf-8-sig')
            except UnicodeDecodeError:
                texto = conteudo.decode('cp1252')
            try:
                dialect = csv.Sniffer().sniff(texto[:8192], delimiters=';,\t')
            except csv.Error:
                dialect = csv.excel
            linhas = csv.reader(io.StringIO(texto), dialect)
        else:
            raise HTTPException(400, 'Use CSV ou XLSX; para XLS, salve como XLSX')
        headers = next(linhas, [])
        campos = [COLUNAS.get(cabecalho(h)) for h in headers]
        conhecidos = [c for c in campos if c]
        if 'telefone' not in conhecidos:
            raise HTTPException(400, 'Inclua a coluna telefone ou whatsapp no cabecalho')
        if len(conhecidos) != len(set(conhecidos)):
            raise HTTPException(400, 'Ha colunas duplicadas para o mesmo campo')
        leads, erros = [], []
        for numero, row in enumerate(linhas, 2):
            if numero > 20001:
                raise HTTPException(400, 'Limite de 20.000 linhas por arquivo')
            if not any(v is not None and str(v).strip() for v in row):
                continue
            lead = {}
            try:
                for campo, valor in zip(campos, row):
                    if not campo or valor is None or valor == '':
                        continue
                    if isinstance(valor, str) and valor.startswith('='):
                        raise ValueError('Substitua formulas por valores')
                    if isinstance(valor, (date, datetime)):
                        valor = valor.strftime('%d/%m/%Y')
                    elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
                        if int(valor) != valor:
                            raise ValueError(f'{campo}: numero fracionado invalido')
                        valor = str(int(valor))
                        if campo == 'cpf':
                            valor = valor.zfill(11)
                    lead[campo] = str(valor).strip()
                if 'variaveis_template' in lead:
                    lead['variaveis_template'] = json.loads(lead['variaveis_template'])
                else:
                    indices = [i for i in range(1, 11) if f'variavel_{i}' in lead]
                    if indices:
                        lead['variaveis_template'] = [lead.pop(f'variavel_{i}', '') for i in range(1, max(indices) + 1)]
                leads.append({'linha': numero, 'lead': lead})
            except (ValueError, TypeError) as e:
                erros.append({'linha': numero, 'motivo': str(e)[:150]})
        return leads, erros
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, 'Arquivo invalido. Confira o formato e o cabecalho')
    finally:
        if workbook:
            workbook.close()
